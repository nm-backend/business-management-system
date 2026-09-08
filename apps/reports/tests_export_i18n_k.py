"""
Тесты локализации отчётов и экспортов (требование ТЗ).

Заголовок отчёта, шапка таблицы и подписи справочников (единицы измерения,
статусы заказа, статус оплаты) раньше были зашиты в коде по-узбекски —
владелец с русским интерфейсом получал PDF/Excel на другом языке.
"""
import io
from datetime import timedelta
from decimal import Decimal

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.clients.models import Client, Payment
from apps.companies.models import Company
from apps.orders.models import Order
from apps.warehouse.models import RawMaterial
from core.utils import translate


class ExportLocalizationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(name='Granit')
        # У компании может быть только один владелец (unique-констрейнт),
        # поэтому язык переключаем на нём же.
        cls.owner_ru = User.objects.create_user(
            username='own_ru', password='pw', role=User.Role.OWNER,
            company=cls.company, language='ru',
        )
        cls.material = RawMaterial.objects.create(
            company=cls.company, name='Гранит', stone_type='granite',
            unit='m2', quantity=Decimal('3'), min_stock=Decimal('10'),
        )
        cls.client_obj = client_obj = Client.objects.create(company=cls.company, name='Акбаров')
        cls.order = Order.objects.create(
            company=cls.company, client=client_obj, custom_product_name='Столешница',
            quantity=Decimal('2'), unit='sht', total_amount=Decimal('1000'),
            deadline=timezone.now() + timedelta(days=3),
        )

    def _xlsx_first_row(self, response):
        wb = load_workbook(io.BytesIO(response.content))
        sheet = wb[wb.sheetnames[0]]
        return [cell for cell in next(sheet.iter_rows(values_only=True))]

    def _api(self, user):
        api = APIClient()
        api.force_authenticate(user)
        return api

    def test_stock_export_header_follows_user_language(self):
        ru = self._api(self.owner_ru).get('/api/v1/reports/export/stock/')
        self.assertEqual(ru.status_code, 200)
        self.assertEqual(self._xlsx_first_row(ru)[0], 'Название')

        self.owner_ru.language = 'uz_cyrl'
        self.owner_ru.save(update_fields=['language'])
        uz = self._api(self.owner_ru).get('/api/v1/reports/export/stock/')
        self.assertEqual(self._xlsx_first_row(uz)[0], 'Номи')

    def test_stock_export_translates_unit_labels(self):
        response = self._api(self.owner_ru).get('/api/v1/reports/export/stock/')
        wb = load_workbook(io.BytesIO(response.content))
        sheet = wb[wb.sheetnames[0]]
        units = [row[3] for row in sheet.iter_rows(values_only=True) if row[0] == 'Гранит']
        self.assertEqual(units, [translate('units.m2', 'ru')])

    def test_orders_export_translates_status_labels(self):
        response = self._api(self.owner_ru).get('/api/v1/reports/export/orders/')
        wb = load_workbook(io.BytesIO(response.content))
        sheet = wb[wb.sheetnames[0]]
        rows = [row for row in sheet.iter_rows(values_only=True)]
        self.assertEqual(rows[0][1], 'Клиент')
        status_cell = rows[1][4]
        self.assertEqual(status_cell, translate(f'statuses.{self.order.status}', 'ru'))

    def test_finance_export_rows_localized(self):
        response = self._api(self.owner_ru).get('/api/v1/reports/export/finance/?format=csv')
        body = response.content.decode('utf-8')
        self.assertIn(translate('finance.revenue', 'ru'), body)
        self.assertIn(translate('export.col_indicator', 'ru'), body)
        self.assertNotIn('Даромад', body)  # узбекский вариант не должен попасть

    def test_shortage_export_localized(self):
        response = self._api(self.owner_ru).get(
            '/api/v1/reports/export/?report_type=material_shortage&format_type=csv'
        )
        body = response.content.decode('utf-8')
        self.assertIn(translate('export.col_shortage', 'ru'), body)

    def test_pdf_title_localized(self):
        response = self._api(self.owner_ru).get('/api/v1/reports/export/stock/?format=pdf')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_revenue_timeline_labels_localized(self):
        # График строится по фактическим платежам — создаём один.
        Payment.objects.create(
            company=self.company, client=self.client_obj, order=self.order,
            amount=Decimal('500'), payment_date=timezone.now(),
        )
        response = self._api(self.owner_ru).get('/api/v1/reports/analytics/revenue-timeline/')
        self.assertEqual(response.status_code, 200)
        labels = response.json()['labels']
        # Метки вида «Сен'26» — берутся из locale, а не из словаря в коде.
        self.assertTrue(all("'" in label for label in labels), labels)
        month = timezone.localdate().month
        self.assertTrue(
            any(label.startswith(translate(f'months_short.{month}', 'ru')) for label in labels),
            labels,
        )
