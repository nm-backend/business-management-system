"""
Выгрузка данных СВОЕЙ компании (tenant-scoped export).

Это отдельный механизм, а не платформенный backup: тот делает pg_dump всей
базы и остаётся у супер-администратора. Здесь владелец забирает только
собственные записи.

Проверяется всё, что делает такую выгрузку опасной: чужой tenant, подмена
company_id, доступ ролей, содержимое файла и след в журнале.
"""
import io
import threading
from decimal import Decimal

from django.db import connection, connections
from django.test import TestCase, TransactionTestCase, tag
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.audit.models import AuditLog
from apps.clients.models import Client, Payment
from apps.companies.models import Company
from apps.finance.models import Expense, ExpenseCategory
from apps.orders.models import Order
from apps.warehouse.models import FinishedProduct, RawMaterial

EXPORT = '/api/v1/reports/export/company-data/'
PLATFORM_BACKUP = ('/api/v1/backup/config/', '/api/v1/backup/logs/', '/api/v1/backup/trigger/')
IS_POSTGRES = connection.vendor == 'postgresql'


class CompanyExportSetupMixin:
    @classmethod
    def _setup(cls):
        cls.company = Company.objects.create(name='MineCo')
        cls.owner = User.objects.create_user(
            username='ce_owner', password='pw', role=User.Role.OWNER, company=cls.company,
        )
        cls.admin = User.objects.create_user(
            username='ce_admin', password='pw', role=User.Role.ADMIN, company=cls.company,
        )
        cls.worker = User.objects.create_user(
            username='ce_worker', password='pw', role=User.Role.WORKER, company=cls.company,
        )
        cls.client_obj = Client.objects.create(
            company=cls.company, name='Мой клиент', phone='+998901112233',
        )
        cls.order = Order.objects.create(
            company=cls.company, client=cls.client_obj, custom_product_name='Моя столешница',
            quantity=Decimal('2'), unit='sht', total_amount=Decimal('111111.11'),
            deadline=timezone.now() + timezone.timedelta(days=2),
        )
        Payment.objects.create(
            company=cls.company, client=cls.client_obj, order=cls.order,
            amount=Decimal('22222.22'), payment_date=timezone.now(),
        )
        RawMaterial.objects.create(
            company=cls.company, name='Мой гранит', unit='m2', quantity=Decimal('10'),
            purchase_price=Decimal('33333.33'),
        )
        FinishedProduct.objects.create(
            company=cls.company, name='Мой товар', unit='sht', quantity=Decimal('1'),
            sale_price=Decimal('44444.44'),
        )
        Expense.objects.create(
            company=cls.company, category=ExpenseCategory.RENT, amount=Decimal('55555.55'),
            created_by=cls.owner, date=timezone.localdate(),
        )

        # Чужая компания с узнаваемыми данными — их в выгрузке быть не должно.
        cls.other = Company.objects.create(name='TheirCo')
        cls.other_owner = User.objects.create_user(
            username='ce_other', password='pw', role=User.Role.OWNER, company=cls.other,
        )
        cls.other_client = Client.objects.create(company=cls.other, name='ЧУЖОЙ КЛИЕНТ')
        Order.objects.create(
            company=cls.other, client=cls.other_client, custom_product_name='ЧУЖОЙ ЗАКАЗ',
            quantity=Decimal('1'), unit='sht', total_amount=Decimal('999999.99'),
            deadline=timezone.now() + timezone.timedelta(days=1),
        )
        RawMaterial.objects.create(
            company=cls.other, name='ЧУЖОЙ МРАМОР', unit='m2', quantity=Decimal('5'),
            purchase_price=Decimal('888888.88'),
        )

    def api_as(self, user):
        api = APIClient()
        api.force_authenticate(user)
        return api

    def workbook_text(self, response):
        """Все ячейки книги одной строкой — по ней ищем утечки."""
        workbook = load_workbook(io.BytesIO(response.content))
        chunks = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                chunks += [str(cell) for cell in row if cell is not None]
        return ' '.join(chunks)


class CompanyDataExportTests(CompanyExportSetupMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls._setup()

    # ── Содержимое ──────────────────────────────────────────────────────────

    def test_owner_gets_own_data(self):
        response = self.api_as(self.owner).get(EXPORT)
        self.assertEqual(response.status_code, 200)
        text = self.workbook_text(response)
        for expected in ('Мой клиент', 'Моя столешница', 'Мой гранит', 'Мой товар',
                         '111111.11', '22222.22', '33333.33', '44444.44', '55555.55'):
            self.assertIn(expected, text, f'в выгрузке нет {expected}')

    def test_export_has_no_other_tenant_data(self):
        text = self.workbook_text(self.api_as(self.owner).get(EXPORT))
        for foreign in ('ЧУЖОЙ КЛИЕНТ', 'ЧУЖОЙ ЗАКАЗ', 'ЧУЖОЙ МРАМОР',
                        '999999.99', '888888.88'):
            self.assertNotIn(foreign, text, f'в выгрузку попало чужое: {foreign}')

    def test_export_has_multiple_sheets(self):
        response = self.api_as(self.owner).get(EXPORT)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertGreaterEqual(len(workbook.worksheets), 7)

    def test_each_owner_gets_only_his_company(self):
        text = self.workbook_text(self.api_as(self.other_owner).get(EXPORT))
        self.assertIn('ЧУЖОЙ КЛИЕНТ', text)
        self.assertNotIn('Мой клиент', text, 'владелец другой компании увидел наши данные')

    # ── Права ───────────────────────────────────────────────────────────────

    def test_admin_never_gets_financial_export(self):
        self.assertEqual(self.api_as(self.admin).get(EXPORT).status_code, 403)

    def test_worker_denied(self):
        self.assertEqual(self.api_as(self.worker).get(EXPORT).status_code, 403)

    def test_anonymous_denied(self):
        self.assertIn(APIClient().get(EXPORT).status_code, (401, 403))

    def test_platform_backup_still_closed_for_owner(self):
        """Платформенный backup остаётся супер-админским — он не тронут."""
        for url in PLATFORM_BACKUP:
            method = self.api_as(self.owner).post if 'trigger' in url else self.api_as(self.owner).get
            self.assertEqual(method(url).status_code, 403, url)

    # ── Подмена tenant ──────────────────────────────────────────────────────

    def test_company_id_in_query_is_ignored(self):
        for param in ('company', 'company_id', 'tenant', 'tenant_id'):
            text = self.workbook_text(
                self.api_as(self.owner).get(EXPORT, {param: self.other.id}),
            )
            self.assertNotIn('ЧУЖОЙ КЛИЕНТ', text, f'подмена через ?{param} сработала')
            self.assertIn('Мой клиент', text)

    def test_company_id_in_body_is_ignored(self):
        """GET с телом — тоже попытка подмены; результат не должен меняться."""
        api = self.api_as(self.owner)
        response = api.generic(
            'GET', EXPORT, data='{"company": %d}' % self.other.id,
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('ЧУЖОЙ КЛИЕНТ', self.workbook_text(response))

    def test_no_path_parameter_exists(self):
        """Пути вида /export/company-data/<id>/ не существует — подставить нечего."""
        self.assertEqual(self.api_as(self.owner).get(f'{EXPORT}{self.other.id}/').status_code, 404)

    # ── Журнал ──────────────────────────────────────────────────────────────

    def test_export_is_audited(self):
        self.api_as(self.owner).get(EXPORT)
        entry = AuditLog.objects.filter(action=AuditLog.Action.EXPORT).latest('created_at')
        self.assertEqual(entry.actor_id, self.owner.id)
        self.assertEqual(entry.metadata.get('export'), 'company_data')

    def test_audit_entry_not_visible_to_other_tenant(self):
        self.api_as(self.owner).get(EXPORT)
        rows = self.api_as(self.other_owner).get('/api/v1/audit/logs/').data
        rows = rows['results'] if 'results' in rows else rows
        self.assertFalse(
            [r for r in rows if r.get('action') == 'export'],
            'чужая компания видит наш экспорт в журнале',
        )


@tag('postgres')
class CompanyExportConcurrencyTests(CompanyExportSetupMixin, TransactionTestCase):
    """Параллельные выгрузки на PostgreSQL."""

    def setUp(self):
        if not IS_POSTGRES:
            self.skipTest('Проверяется только на PostgreSQL')
        super().setUp()
        self._setup()

    def test_parallel_exports_do_not_mix_tenants(self):
        """
        Владельцы двух компаний выгружают данные одновременно.

        Проверяем, что запросы не перепутали компании: каждый получает своё.
        """
        results = {}
        barrier = threading.Barrier(2)

        def run(user, key):
            try:
                barrier.wait(timeout=30)
                response = self.api_as(user).get(EXPORT)
                results[key] = (response.status_code, self.workbook_text(response))
            finally:
                connections.close_all()

        threads = [
            threading.Thread(target=run, args=(self.owner, 'mine')),
            threading.Thread(target=run, args=(self.other_owner, 'theirs')),
        ]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=60)

        self.assertEqual(results['mine'][0], 200)
        self.assertEqual(results['theirs'][0], 200)
        self.assertIn('Мой клиент', results['mine'][1])
        self.assertNotIn('ЧУЖОЙ КЛИЕНТ', results['mine'][1])
        self.assertIn('ЧУЖОЙ КЛИЕНТ', results['theirs'][1])
        self.assertNotIn('Мой клиент', results['theirs'][1])
