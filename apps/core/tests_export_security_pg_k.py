"""
Adversarial-аудит на PostgreSQL: экспорты, JWT, рецепты, просрочка, изоляция.

Отличие от tests_tz_compliance_k: там проверялись JSON-ответы, здесь —
БИНАРНЫЕ выгрузки (PDF/Excel) и жизненный цикл токенов. Утечка через файл
опаснее: JSON смотрят глазами в браузере, а выгрузку пересылают дальше.

PDF разбирается извлечением текста (pypdf), Excel — по ячейкам (openpyxl):
проверять «нет ли числа в байтах файла» бессмысленно, оно может быть сжато.
"""
import io
import re
from decimal import Decimal

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.clients.models import Client, Payment
from apps.companies.models import Company
from apps.finance.models import Expense, ExpenseCategory, LaborRate
from apps.orders.models import Order
from apps.production.models import Task, TaskStatus, WorkRecord
from apps.warehouse.models import FinishedProduct, RawMaterial, Recipe, RecipeItem


def pdf_text(content: bytes) -> str:
    """Текст всех страниц PDF — по нему и ищем утечки."""
    reader = PdfReader(io.BytesIO(content))
    return '\n'.join(page.extract_text() or '' for page in reader.pages)


def xlsx_values(content: bytes) -> list[str]:
    """Все непустые ячейки книги как строки."""
    workbook = load_workbook(io.BytesIO(content))
    values = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values += [str(cell) for cell in row if cell is not None]
    return values


class ExportScanMixin:
    """Компания, где у каждой сущности своя узнаваемая сумма."""

    # Суммы-«маячки»: если хоть одна попала в выгрузку роли без прав — утечка.
    SECRET_AMOUNTS = {
        'purchase_price': '777.11',
        'cost_price': '888.22',
        'sale_price': '999.33',
        'order_total': '123456.78',
        'payment': '654.32',
        'expense': '4321.09',
        'labor_rate': '345.67',
    }

    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(name='ExportCo')
        cls.owner = User.objects.create_user(
            username='exp_owner', password='pw', role=User.Role.OWNER, company=cls.company,
        )
        cls.admin = User.objects.create_user(
            username='exp_admin', password='pw', role=User.Role.ADMIN, company=cls.company,
        )
        cls.worker = User.objects.create_user(
            username='exp_worker', password='pw', role=User.Role.WORKER,
            company=cls.company, full_name='Али',
        )
        cls.material = RawMaterial.objects.create(
            company=cls.company, name='Гранит', unit='m2', quantity=Decimal('50'),
            min_stock=Decimal('5'),
            purchase_price=Decimal(cls.SECRET_AMOUNTS['purchase_price']),
            avg_cost_price=Decimal(cls.SECRET_AMOUNTS['purchase_price']),
        )
        cls.product = FinishedProduct.objects.create(
            company=cls.company, name='Столешница', unit='sht', quantity=Decimal('3'),
            cost_price=Decimal(cls.SECRET_AMOUNTS['cost_price']),
            sale_price=Decimal(cls.SECRET_AMOUNTS['sale_price']),
        )
        cls.client_obj = Client.objects.create(company=cls.company, name='Акбаров')
        cls.order = Order.objects.create(
            company=cls.company, client=cls.client_obj, product=cls.product,
            quantity=Decimal('1'), unit='sht',
            total_amount=Decimal(cls.SECRET_AMOUNTS['order_total']),
            deadline=timezone.now() + timezone.timedelta(days=2),
        )
        Payment.objects.create(
            company=cls.company, client=cls.client_obj, order=cls.order,
            amount=Decimal(cls.SECRET_AMOUNTS['payment']), payment_date=timezone.now(),
        )
        Expense.objects.create(
            company=cls.company, category=ExpenseCategory.RENT,
            amount=Decimal(cls.SECRET_AMOUNTS['expense']), created_by=cls.owner,
            date=timezone.now().date(),
        )
        LaborRate.objects.create(
            company=cls.company, product=cls.product,
            operation=LaborRate.OperationType.POLISHING,
            rate_per_unit=Decimal(cls.SECRET_AMOUNTS['labor_rate']), unit='sht',
        )
        WorkRecord.objects.create(
            company=cls.company, worker=cls.worker, product=cls.product,
            quantity=Decimal('2'), unit='sht', status=WorkRecord.WorkStatus.CONFIRMED,
            labor_cost=Decimal('691.34'),
        )

    def api_as(self, user):
        api = APIClient()
        api.force_authenticate(user)
        return api

    def assert_no_secret_amounts(self, haystack, source):
        """
        Ни одна «секретная» сумма не должна встречаться в выгрузке.

        Сравниваем и с точкой, и с запятой, и без разделителя тысяч — форматы
        в PDF/Excel разные, а утечка от формата не зависит.
        """
        found = []
        for name, amount in self.SECRET_AMOUNTS.items():
            variants = {
                amount,
                amount.replace('.', ','),
                amount.split('.')[0],                      # 123456
                f'{int(float(amount)):,}'.replace(',', ' '),  # 123 456
            }
            for variant in variants:
                if len(variant) >= 4 and variant in haystack:
                    found.append(f'{name}={variant}')
                    break
        self.assertFalse(found, f'{source}: в выгрузке найдены суммы {found}')


class AdminExportLeakTests(ExportScanMixin, TestCase):
    """Выгрузки администратора не содержат ни одной суммы — включая PDF и Excel."""

    ADMIN_EXPORTS = [
        '/api/v1/reports/export/stock/',
        '/api/v1/reports/export/orders/',
        '/api/v1/reports/export/work/',
        '/api/v1/reports/export/?report_type=material_shortage',
    ]

    def test_admin_xlsx_exports_clean(self):
        api = self.api_as(self.admin)
        for endpoint in self.ADMIN_EXPORTS:
            response = api.get(endpoint)
            self.assertEqual(response.status_code, 200, endpoint)
            self.assert_no_secret_amounts(' '.join(xlsx_values(response.content)), f'XLSX {endpoint}')

    def test_admin_pdf_exports_clean(self):
        api = self.api_as(self.admin)
        for endpoint in self.ADMIN_EXPORTS:
            sep = '&' if '?' in endpoint else '?'
            param = 'format_type=pdf' if 'report_type' in endpoint else 'format=pdf'
            response = api.get(f'{endpoint}{sep}{param}')
            self.assertEqual(response.status_code, 200, endpoint)
            self.assertTrue(response.content.startswith(b'%PDF'), endpoint)
            self.assert_no_secret_amounts(pdf_text(response.content), f'PDF {endpoint}')

    def test_admin_csv_exports_clean(self):
        api = self.api_as(self.admin)
        for endpoint in self.ADMIN_EXPORTS:
            sep = '&' if '?' in endpoint else '?'
            param = 'format_type=csv' if 'report_type' in endpoint else 'format=csv'
            response = api.get(f'{endpoint}{sep}{param}')
            self.assertEqual(response.status_code, 200, endpoint)
            self.assert_no_secret_amounts(response.content.decode('utf-8'), f'CSV {endpoint}')

    def test_admin_cannot_download_finance_export(self):
        api = self.api_as(self.admin)
        for fmt in ('xlsx', 'pdf', 'csv'):
            response = api.get(f'/api/v1/reports/export/finance/?format={fmt}')
            self.assertEqual(response.status_code, 403, fmt)

    def test_worker_cannot_download_any_export(self):
        api = self.api_as(self.worker)
        for endpoint in self.ADMIN_EXPORTS + ['/api/v1/reports/export/finance/']:
            response = api.get(endpoint)
            self.assertEqual(response.status_code, 403, f'{endpoint} -> {response.status_code}')

    def test_owner_export_does_contain_money(self):
        """Контроль осмысленности: у владельца суммы в выгрузке ЕСТЬ."""
        api = self.api_as(self.owner)
        response = api.get('/api/v1/reports/export/finance/?format=csv')
        self.assertEqual(response.status_code, 200)
        body = response.content.decode('utf-8')
        self.assertIn('654.32', body.replace(' ', ''))  # оплата попала в выручку

    def test_pdf_scanner_is_not_blind(self):
        """
        Контроль самого метода проверки: в PDF владельца сумма ДОЛЖНА находиться.

        Без этого теста «чистые» админские PDF ничего не доказывают: пустой
        результат извлечения текста выглядел бы как отсутствие утечки.
        """
        api = self.api_as(self.owner)
        response = api.get('/api/v1/reports/export/finance/?format=pdf')
        self.assertEqual(response.status_code, 200)
        text = pdf_text(response.content)
        self.assertGreater(len(text), 50, 'из PDF не извлёкся текст — сканер слеп')
        self.assertIn('654.32', text.replace(' ', ''))


class JwtLifecycleTests(TestCase):
    """Вход, обновление, выход и отзыв токена."""

    @classmethod
    def setUpTestData(cls):
        cls.company = Company.objects.create(name='JwtCo')
        cls.owner = User.objects.create_user(
            username='jwt_owner', password='Str0ng!Pass9', role=User.Role.OWNER,
            company=cls.company,
        )

    def login(self):
        response = APIClient().post('/api/v1/accounts/login/', {
            'username': 'jwt_owner', 'password': 'Str0ng!Pass9',
        }, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        return response.data['tokens']

    def test_login_returns_pair_and_access_works(self):
        tokens = self.login()
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION=f'Bearer {tokens["access"]}')
        self.assertEqual(api.get('/api/v1/accounts/me/').status_code, 200)

    def test_refresh_issues_new_access(self):
        tokens = self.login()
        response = APIClient().post('/api/v1/accounts/token/refresh/', {
            'refresh': tokens['refresh'],
        }, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn('access', response.data)

    def test_logout_revokes_refresh_token(self):
        """После выхода refresh не должен выдавать новый доступ."""
        tokens = self.login()
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION=f'Bearer {tokens["access"]}')
        logout = api.post('/api/v1/accounts/logout/', {'refresh': tokens['refresh']}, format='json')
        self.assertIn(logout.status_code, (200, 204, 205), logout.data)

        again = APIClient().post('/api/v1/accounts/token/refresh/', {
            'refresh': tokens['refresh'],
        }, format='json')
        self.assertEqual(again.status_code, 401, 'refresh после выхода обязан быть отозван')

    def test_garbage_token_rejected(self):
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION='Bearer not.a.token')
        self.assertEqual(api.get('/api/v1/accounts/me/').status_code, 401)

    def test_inactive_user_cannot_login(self):
        self.owner.is_active = False
        self.owner.save(update_fields=['is_active'])
        response = APIClient().post('/api/v1/accounts/login/', {
            'username': 'jwt_owner', 'password': 'Str0ng!Pass9',
        }, format='json')
        self.assertIn(response.status_code, (400, 401, 403))


class RecipeAndShortageTests(ExportScanMixin, TestCase):
    """Рецепт: количество, единицы, списание и нехватка."""

    def setUp(self):
        self.api = self.api_as(self.owner)
        self.recipe = Recipe.objects.create(
            company=self.company, product=self.product, name='Столешница',
        )
        RecipeItem.objects.create(
            recipe=self.recipe, material=self.material,
            quantity_required=Decimal('2.5'), unit='m2',
        )

    def _work(self, quantity):
        task = Task.objects.create(
            company=self.company, worker=self.worker, assigned_by=self.owner,
            status=TaskStatus.ACCEPTED, title='Столешница',
        )
        return WorkRecord.objects.create(
            company=self.company, task=task, worker=self.worker, product=self.product,
            quantity=Decimal(quantity), unit='sht',
            status=WorkRecord.WorkStatus.AWAITING_CONFIRMATION,
        )

    def test_fractional_recipe_quantity_written_off_exactly(self):
        """2.5 м² × 4 изделия = 10 м² ровно, без потери копеек на округлении."""
        self.material.quantity = Decimal('20')
        self.material.save(update_fields=['quantity'])
        work = self._work('4')

        response = self.api.post(f'/api/v1/production/works/{work.id}/confirm/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)

        self.material.refresh_from_db()
        self.assertEqual(self.material.quantity, Decimal('10.000'))

    def test_defect_consumes_material_too(self):
        """Сырьё уходит и на брак: иначе остаток завышен."""
        self.material.quantity = Decimal('20')
        self.material.save(update_fields=['quantity'])
        work = self._work('2')
        work.defect_quantity = Decimal('1')
        work.save(update_fields=['defect_quantity'])

        response = self.api.post(f'/api/v1/production/works/{work.id}/confirm/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)

        self.material.refresh_from_db()
        # (2 годных + 1 брак) × 2.5 = 7.5
        self.assertEqual(self.material.quantity, Decimal('12.500'))
        self.product.refresh_from_db()
        # На склад попадает только годное.
        self.assertEqual(self.product.quantity, Decimal('5.000'))

    def test_shortage_blocks_confirmation_with_details(self):
        self.material.quantity = Decimal('1')
        self.material.save(update_fields=['quantity'])
        work = self._work('2')

        response = self.api.post(f'/api/v1/production/works/{work.id}/confirm/', {}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        body = str(response.data)
        self.assertIn('Гранит', body, 'в ошибке должен быть виден недостающий материал')

        self.material.refresh_from_db()
        self.assertEqual(self.material.quantity, Decimal('1.000'), 'склад не должен меняться')

    def test_recipe_item_unit_recorded(self):
        item = RecipeItem.objects.get(recipe=self.recipe)
        self.assertEqual(item.unit, 'm2')
        self.assertEqual(item.unit, self.material.unit,
                         'единица позиции рецепта должна совпадать с единицей материала')


class OverdueDebtTests(ExportScanMixin, TestCase):
    """Просрочка долга: срок заказа + неполная оплата."""

    def test_overdue_notification_created_once(self):
        from django.core.management import call_command

        from apps.messaging.models import Notification

        self.order.deadline = timezone.now() - timezone.timedelta(days=3)
        self.order.paid_amount = Decimal('0')
        self.order.save(update_fields=['deadline', 'paid_amount'])

        call_command('notify_overdue_debts')
        first = Notification.objects.filter(
            type=Notification.NotificationType.OVERDUE_DEBT, related_order=self.order,
        ).count()
        self.assertGreater(first, 0, 'просроченный долг должен уведомлять')

        call_command('notify_overdue_debts')
        second = Notification.objects.filter(
            type=Notification.NotificationType.OVERDUE_DEBT, related_order=self.order,
        ).count()
        self.assertEqual(first, second, 'повторный запуск не должен дублировать уведомления')

    def test_paid_order_is_not_overdue(self):
        from django.core.management import call_command

        from apps.messaging.models import Notification

        self.order.deadline = timezone.now() - timezone.timedelta(days=3)
        self.order.paid_amount = self.order.total_amount
        self.order.save(update_fields=['deadline', 'paid_amount'])

        call_command('notify_overdue_debts')
        self.assertEqual(Notification.objects.filter(
            type=Notification.NotificationType.OVERDUE_DEBT, related_order=self.order,
        ).count(), 0)

    def test_overdue_notification_has_no_amount(self):
        """Уведомление о просрочке уходит и администратору — сумм в нём быть не должно."""
        from django.core.management import call_command

        from apps.messaging.models import Notification

        self.order.deadline = timezone.now() - timezone.timedelta(days=1)
        self.order.paid_amount = Decimal('0')
        self.order.save(update_fields=['deadline', 'paid_amount'])
        call_command('notify_overdue_debts')

        for notification in Notification.objects.filter(
            type=Notification.NotificationType.OVERDUE_DEBT,
        ):
            self.assertNotIn(self.SECRET_AMOUNTS['order_total'], notification.message)
            self.assertFalse(
                re.search(r'\d{4,}', notification.message.replace(f'#{self.order.id}', '')),
                f'в тексте просрочки видна сумма: {notification.message}',
            )


class SupplierFieldTests(ExportScanMixin, TestCase):
    """Поставщик: сейчас это свободный текст — фиксируем фактическое поведение."""

    def test_supplier_is_free_text_and_saved(self):
        api = self.api_as(self.owner)
        response = api.post('/api/v1/warehouse/raw-materials/', {
            'name': 'Мрамор', 'unit': 'm2', 'supplier': 'ООО «Тошкент Мармар»',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['supplier'], 'ООО «Тошкент Мармар»')

    def test_supplier_visible_to_admin_without_prices(self):
        """Поставщик — операционные данные: админу виден, цены рядом — нет."""
        self.material.supplier = 'ООО «Тошкент Мармар»'
        self.material.save(update_fields=['supplier'])
        api = self.api_as(self.admin)
        rows = api.get('/api/v1/warehouse/raw-materials/').data
        rows = rows['results'] if 'results' in rows else rows
        row = next(r for r in rows if r['name'] == 'Гранит')
        self.assertEqual(row['supplier'], 'ООО «Тошкент Мармар»')
        self.assertNotIn('purchase_price', row)
